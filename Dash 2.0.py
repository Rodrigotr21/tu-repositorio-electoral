import streamlit as st
import pandas as pd
import os
import io
import datetime  # <-- Para generar nombres únicos de archivos
import base64    # <-- Para procesar y enviar las imágenes a Apps Script
import requests  # <-- Para conectarnos con tu Google Drive a través de tu URL

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Inicialización de estados de Streamlit para control de borrado automático y mensajes persistentes
if "widget_counter" not in st.session_state:
    st.session_state.widget_counter = 0
if "mensaje_exito" not in st.session_state:
    st.session_state.mensaje_exito = None

# Configuración de la página
st.set_page_config(
    page_title="Gestión de Formatos Electorales | GOECOR",
    page_icon="🗳️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Estilos personalizados para mejorar la apariencia
st.markdown("""
    <style>
    .main-title {
        color: #1F4E79;
        font-size: 32px;
        font-weight: bold;
        margin-bottom: 5px;
    }
    .subtitle {
        color: #595959;
        font-size: 16px;
        margin-bottom: 25px;
    }
    .card {
        background-color: #F8FAFC;
        padding: 20px;
        border-radius: 10px;
        border-left: 5px solid #1F4E79;
        margin-bottom: 20px;
    }
    .signature-box {
        background-color: #F1F5F9;
        padding: 15px;
        border-radius: 5px;
        border: 1px dashed #94A3B8;
        font-weight: bold;
        color: #334155;
    }
    </style>
""", unsafe_allow_html=True)

# Base de datos estructurada - JEL
FORMATOS_JEL = {
    "FM02": {
        "codigo": "FM02-GOECOR/JEL",
        "nombre": "Lista de Chequeo de Acondicionamiento y Señalización del Local de Votación",
        "proposito": "Verificar el estado del acondicionamiento físico, accesos y la señalización del local de votación, asegurando el estricto cumplimiento de las normativas vigentes estipuladas por la GOECOR.",
        "firmas": "Coordinador de Local de Votación (CLV)",
        "tipo_descarga": "Word (Checklist)",
        "columnas": ["N°", "Aspecto a Evaluar / Elemento de Señalización", "Cumple (Sí)", "No Cumple (No)", "No Aplica (N/A)", "Observaciones / Detalles"],
        "datos_ejemplo": [
            ["1", "Señalización externa e interna de accesos, ingreso y salida clara", "", "", "", ""],
            ["2", "Rampas y vías de acceso libres de obstáculos para personas con discapacidad", "", "", "", ""],
            ["3", "Acondicionamiento de aulas y ubicación de mesas según croquis de la directiva", "", "", "", ""],
            ["4", "Publicación visible de la lista de electores en el exterior del local", "", "", "", ""],
            ["5", "Iluminación adecuada y operativa en pasadizos, escaleras y aulas de sufragio", "", "", "", ""]
        ]
    },
    "FM03": {
        "codigo": "FM03-GOECOR/JEL",
        "nombre": "Control de Distribución del Material Electoral durante la Jornada Electoral",
        "proposito": "Controlar el reparto detallado del material electoral y refrigerios desde el centro de acopio local hacia las mesas de sufragio.",
        "firmas": "No requiere firma (Documento operativo de control interno)",
        "tipo_descarga": "Excel (Control Operativo)",
        "columnas": [
            "N°", "Nombre del CM o CM STAE", "N° Mesa de Sufragio", "Hora Distribución (HH:MM)", 
            "Caja Material (*)", "Caja Refrigerio (*)", "Presidente (Titular)", "Secretario (Titular)", 
            "3er Miembro (Titular)", "1er Suplente", "2do Suplente", "Electores de la Fila (*)"
        ],
        "datos_ejemplo": [
            ["1", "Pérez Gómez Juan Carlos", "012345", "06:15", "✓", "✓", "✓", "✓", "✓", "", "", ""],
            ["2", "Rojas Mendoza María Fe", "012346", "06:22", "✓", "✓", "✓", "✓", "", "✓", "", ""]
        ]
    },
    "FM05": {
        "codigo": "FM05-GOECOR/JEL",
        "nombre": "Reporte de la Jornada Electoral (SIDE en Físico)",
        "proposito": "Registrar de forma física las variables críticas de la jornada electoral para su transmisión al Centro de Soporte.",
        "firmas": "Sin firmas (Reporte informativo técnico para transmisión)",
        "tipo_descarga": "Word (Reporte Consolidado)",
        "columnas": ["Sección / Numeral", "Variable Electoral / Descripción del Reporte", "Valor / Registro Requerido", "Instrucciones de Llenado"],
        "datos_ejemplo": [
            ["1", "CANTIDAD DE MESAS INSTALADAS", "______", "Indicar el total de mesas operativas en el local."],
            ["2", "NÚMERO Y HORA DE INSTALACIÓN DE 1ra MESA DE SUFRAGIO", "Mesa: _________ / Hora: ____:____", "Reportar inmediatamente al Centro de Soporte."]
        ]
    },
    "FM07": {
        "codigo": "FM07-GOECOR/JEL",
        "nombre": "Control de Documentos y Materiales Electorales en el Centro de Acopio",
        "proposito": "Gestionar y validar rigurosamente la recepción y repliegue de las actas y sobres electorales de seguridad devueltos.",
        "firmas": "Coordinador de Local de Votación (CLV)",
        "tipo_descarga": "Excel (Validación de Actas)",
        "columnas": [
            "N°", "N° de Mesa de Sufragio", "Sobre Plomo (√)", "Sobre Verde (√)", 
            "Sobre Celeste (√)", "Sobre Rojo (√)", "Caja Restos Electorales (√)", 
            "Sobre Anaranjado (√)", "Sobre Cédulas No Impugnadas (√)"
        ],
        "datos_ejemplo": [
            ["1", "012345", "√", "√", "√", "√", "√", "√", "√"],
            ["2", "012346", "√", "√", "√", "√", "√", "√", "√"]
        ]
    },
    "FM11": {
        "codigo": "FM11-GOECOR/JEL",
        "nombre": "Acta de Recepción y Devolución del Local de Votación",
        "proposito": "Documentar formalmente el estado de la infraestructura civil y aulas cedidas por la institución educativa.",
        "firmas": "Coordinador de Local de Votación y Responsable Designado",
        "tipo_descarga": "Word (Acta Formal)",
        "columnas": ["Condiciones de Recepción/Devolución de Ambientes", "Especificación de Aula(s) / Ubicación", "Estado en Recepción (PRE)", "Estado en Devolución (POST)"],
        "datos_ejemplo": [
            ["Ambientes Cedidos (Aulas / Patios)", "Aulas del 101 al 115, Patio Central", "Recibido Conforme", "Devuelto Conforme"],
            ["Ventanas rotas / Vidrios dañados", "Aula 104", "01 vidrio rajado al ingreso", "Permanece en igual condición"]
        ]
    },
    "FM12": {
        "codigo": "FM12-GOECOR/JEL",
        "nombre": "Hoja de Ruta para el Traslado de Sobres Plomos por Entregas a la Sede ODPE",
        "proposito": "Garantizar la trazabilidad absoluta en el traslado físico de las actas electorales (sobres plomos).",
        "firmas": "Responsable del Traslado de Sobres, Coordinador de Local de Votación y Encargado ODPE",
        "tipo_descarga": "Excel (Hoja de Ruta)",
        "columnas": ["N°", "N° de la Mesa de Sufragio", "Elección / Consulta 1 (✓)", "Elección / Consulta 2 (✓)", "Elección / Consulta 3 (✓)", "Firma del Responsable del Traslado"],
        "datos_ejemplo": [
            ["1", "012345", "✓", "", "", "___________________________"]
        ]
    },
    "FM13": {
        "codigo": "FM13-GOECOR/JEL",
        "nombre": "Hoja de Ruta para el Traslado de Documentos Electorales desde el Local hacia la ODPE",
        "proposito": "Monitorear el repliegue seguro y completo de los documentos electorales restantes y sobres con cédulas.",
        "firmas": "Responsable del Traslado de Sobres, Coordinador de Local de Votación y Encargado ODPE",
        "tipo_descarga": "Excel (Trazabilidad)",
        "columnas": ["N°", "N° Mesa de Sufragio", "Sobre con Cédulas de Sufragio (✓)", "Elección / Consulta 1 (✓)", "Elección / Consulta 2 (✓)", "Elección / Consulta 3 (✓)"],
        "datos_ejemplo": [
            ["1", "012345", "✓", "✓", "✓", "✓"]
        ]
    },
    "FM15": {
        "codigo": "FM15-GOECOR/JEL",
        "nombre": "Cargo de Entrega del Material de Reserva al CM / CM STAE",
        "proposito": "Registrar de manera transparente cualquier asignación excepcional de material electoral de contingencia o reserva.",
        "firmas": "Coordinador de Local de Votación (CLV) y Coordinador de Mesa (CM / CM STAE) receptor",
        "tipo_descarga": "Word (Cargo de Entrega)",
        "columnas": ["Descripción del Material de Reserva", "Cantidad de Material Entregado", "Número de Mesa de Sufragio", "Nombres y Apellidos del CM / CM STAE", "Firma del CM / CM STAE"],
        "datos_ejemplo": [
            ["Tinta Indelebre de contingencia", "1 unidad", "012345", "Pedro Alcántara Gómez", "[Firma Física]"]
        ]
    },
    "FM16": {
        "codigo": "FM16-GOECOR/JEL",
        "nombre": "Cargo de Recepción de Sobres Verdes, Celestes, Rojos y Anaranjados en Custodia",
        "proposito": "Contabilizar y custodiar en forma matricial los sobres de seguridad de colores.",
        "firmas": "Responsable de Local de Votación (RLV), Responsable del Traslado de Sobres y Encargado ODPE",
        "tipo_descarga": "Word (Matriz de Control)",
        "columnas": ["Tipo de Sobre", "Elección / Consulta 1 (Ánforas)", "Elección / Consulta 2 (Ánforas)", "Elección / Consulta 3 (Ánforas)", "Total General Custodiado"],
        "datos_ejemplo": [
            ["SOBRES ROJOS", "15", "0", "0", "15 Ánforas"],
            ["SOBRES VERDES", "15", "15", "0", "30 Ánforas"]
        ]
    }
}

# Base de datos estructurada - CAE
FORMATOS_CAE = {
    "FM01": {
        "codigo": "FM01-GOECOR/CAE",
        "nombre": "Registro de Capacitación de Miembros de Mesa a Domicilio y en Línea",
        "proposito": "Registrar de forma oficial a los miembros de mesa que reciben capacitación electoral.",
        "firmas": "Coordinador / Personal encargado de realizar la capacitación",
        "tipo_descarga": "Excel (Registro Integral)",
        "columnas": ["N°", "Nombres y Apellidos", "DNI", "N.° de Mesa", "Cargo", "Fecha (dd/mm/aa)", "Estrategia", "Teléfono", "Correo Electrónico", "Firma del Miembro de Mesa", "Reportado al SIGCAE"],
        "datos_ejemplo": [
            ["1", "Alva Prado Manuel Enrique", "45678912", "045123", "Presidente", "12/05/2026", "A domicilio", "999888777", "m.alva@mail.com", "[Firma Física]", "✓"]
        ]
    },
    "FM02": {
        "codigo": "FM02-GOECOR/CAE",
        "nombre": "Registro de Capacitación de Electores",
        "proposito": "Controlar y registrar a los ciudadanos electores capacitados.",
        "firmas": "Capacitador / Personal de la Oficina Distrital o de Coordinación",
        "tipo_descarga": "Excel (Control de Electores)",
        "columnas": ["N°", "Nombres y Apellidos", "DNI", "Fecha (dd/mm/aa)", "Estrategia", "Teléfono", "Correo Electrónico", "Firma / Palote", "Reportado al SIGCAE"],
        "datos_ejemplo": [
            ["1", "Bendezú Palomino Jorge Luis", "10234567", "14/05/2026", "Personalizada", "977666555", "j.bendezu@mail.com", "[Firma]", "✓"]
        ]
    },
    "FM03": {
        "codigo": "FM03-GOECOR/CAE",
        "nombre": "Registro de Capacitación de Miembros de Mesa en Áreas de Capacitación",
        "proposito": "Registrar a cada miembro de mesa capacitado de forma presencial en las áreas fijas.",
        "firmas": "Capacitador / Especialista de Capacitación Electoral",
        "tipo_descarga": "Excel (Áreas Fijas)",
        "columnas": ["N°", "Nombres y Apellidos", "DNI", "N.º de Mesa de Sufragio", "Cargo", "Estrategia", "Fecha (dd/mm/aa)", "Teléfono", "Correo Electrónico", "Firma", "Reportado al SIGCAE"],
        "datos_ejemplo": [
            ["1", "Gómez Peralta Ricardo", "40567891", "012345", "3er Miembro", "Taller", "15/05/2026", "955444333", "r.gomez@mail.com", "[Firma]", "✓"]
        ]
    },
    "FM04": {
        "codigo": "FM04-GOECOR/CAE",
        "nombre": "Registro de Capacitación de Personeros",
        "proposito": "Registrar formalmente la asistencia de los personeros de las diversas agrupaciones políticas.",
        "firmas": "Personal Responsable de la Capacitación de Personeros",
        "tipo_descarga": "Excel (Control Político)",
        "columnas": ["N°", "Nombres y Apellidos", "DNI", "Nombre de la Organización Política", "Estrategia", "Fecha (dd/mm/aa)", "Teléfono", "Correo Electrónico", "Firma", "Reportado al SIGCAE"],
        "datos_ejemplo": [
            ["1", "Quispe Mamani Wilfredo", "09876543", "Frente Político de Integración", "En reunión", "16/05/2026", "933222111", "w.quispe@mail.com", "[Firma]", "✓"]
        ]
    },
    "FM05": {
        "codigo": "FM05-GOECOR/CAE",
        "nombre": "Registro de Capacitación a Efectivos de las FF. AA. y la PNP",
        "proposito": "Consolidar cuantitativa y nominalmente el número de efectivos militares y policiales instruidos.",
        "firmas": "Responsable de las FF. AA. y Responsable de la PNP asignados",
        "tipo_descarga": "Excel (Consolidado de Fuerzas)",
        "columnas": ["Institución de Resguardo", "Hombres Capacitados", "Mujeres Capacitadas", "Total Efectivos", "Nombre del Jefe / Responsable", "Cargo o Grado Oficial"],
        "datos_ejemplo": [
            ["Fuerzas Armadas (FF. AA.)", "45", "15", "60", "Gral. Carlos Montoya P.", "Comandante de Plaza Local"]
        ]
    },
    "FM06": {
        "codigo": "FM06-GOECOR/CAE",
        "nombre": "Registro de Capacitación de Miembros de Mesa en Jornada de Capacitación por Aula",
        "proposito": "Llevar el control estricto por aulas durante las jornadas nacionales de capacitación masiva.",
        "firmas": "Coordinador de Aula / Capacitador Responsable",
        "tipo_descarga": "Excel (Control de Aula Masivo)",
        "columnas": ["N°", "Nombres y Apellidos", "DNI", "N.º de Mesa de Sufragio", "Cargo", "Teléfono", "Correo Electrónico", "Firma", "Observaciones / Detalles"],
        "datos_ejemplo": [
            ["1", "Ramírez Ortiz Fernando", "44556677", "023456", "Presidente", "911222333", "f.ramirez@mail.com", "[Firma]", "Asistió conforme"]
        ]
    }
}

NOMBRES_EXACTOS_JEL = {
    "FM02": "FM02-GOECOR-JEL.docx",
    "FM03": "FM03-GOECOR-JEL.xlsx",
    "FM05": "FM05-GOECOR-JEL.docx",
    "FM07": "FM07-GOECOR-JEL.xlsx",
    "FM11": "FM11-GOECOR-JEL.docx",
    "FM12": "FM12-GOECOR-JEL.xlsx",
    "FM13": "FM13-GOECOR-JEL.xlsx",
    "FM15": "FM15-GOECOR-JEL.docx",
    "FM16": "FM16-GOECOR-JEL.docx"
}

NOMBRES_EXACTOS_CAE = {
    "FM01": "FM01-GOECOR-CAE.xlsx",
    "FM02": "FM02-GOECOR-CAE.xlsx",
    "FM03": "FM03-GOECOR-CAE.xlsx",
    "FM04": "FM04-GOECOR-CAE.xlsx",
    "FM05": "FM05-GOECOR-CAE.xlsx",
    "FM06": "FM06-GOECOR-CAE.xlsx"
}

def crear_excel_formato(id_formato, info):
    """Generador de Plan B en caso de que falte el archivo físico"""
    wb = Workbook()
    ws = wb.active
    ws.title = id_formato
    
    ws.views.sheetView[0].showGridLines = True
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    
    font_title = Font(name="Arial", size=14, bold=True, color="1F4E79")
    font_subtitle = Font(name="Arial", size=10, italic=True, color="595959")
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Arial", size=10)
    font_bold = Font(name="Arial", size=10, bold=True)
    
    thin_side = Side(border_style="thin", color="D9D9D9")
    thick_bottom = Side(border_style="medium", color="1F4E79")
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    
    ws["A1"] = info["codigo"]
    ws["A1"].font = font_title
    ws.merge_cells("A1:E1")
    
    ws["A2"] = info["nombre"]
    ws["A2"].font = Font(name="Arial", size=12, bold=True, color="333333")
    ws.merge_cells("A2:E2")
    
    ws["A3"] = f"Propósito: {info['proposito']}"
    ws["A3"].font = font_subtitle
    ws["A3"].alignment = align_left
    ws.merge_cells("A3:E3")
    
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 28
    
    start_row = 5
    
    for col_num, header_text in enumerate(info["columnas"], start=1):
        cell = ws.cell(row=start_row, column=col_num, value=header_text)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center if any(x in header_text.lower() for x in ["n°", "hora", "✓", "√", "sí", "no"]) else align_left
        cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thick_bottom)
    ws.row_dimensions[start_row].height = 26
    
    current_row = start_row + 1
    for i, row_data in enumerate(info["datos_ejemplo"]):
        for col_num, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_num, value=val)
            cell.font = font_body
            cell.border = border_cell
            
            if col_num == 1 or col_num == 2:
                cell.alignment = align_left
            else:
                cell.alignment = align_center if len(str(val)) < 15 else align_left
                
            if i % 2 == 1:
                cell.fill = zebra_fill
                
        ws.row_dimensions[current_row].height = 20
        current_row += 1
        
    for _ in range(5):
        for col_num in range(1, len(info["columnas"]) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.border = border_cell
            if (_ + len(info["datos_ejemplo"])) % 2 == 1:
                cell.fill = zebra_fill
        ws.row_dimensions[current_row].height = 20
        current_row += 1
        
    current_row += 2
    ws.cell(row=current_row, column=1, value="CONTROL DE VALIDACIÓN Y FIRMAS REGISTRADAS:").font = font_bold
    current_row += 1
    
    firmas_req = info["firmas"]
    ws.cell(row=current_row, column=1, value=f"Responsables Obligatorios: {firmas_req}").font = font_subtitle
    current_row += 2
    
    if "No requiere" not in firmas_req and "Sin firmas" not in firmas_req:
        lista_firmas = [f.strip() for f in firmas_req.replace(" y ", ",").split(",")]
        for idx, f_name in enumerate(lista_firmas):
            col_pos = (idx * 2) + 1
            if col_pos <= len(info["columnas"]):
                ws.cell(row=current_row, column=col_pos, value="___________________________").alignment = align_center
                ws.cell(row=current_row+1, column=col_pos, value=f_name).font = Font(name="Arial", size=9, bold=True)
                ws.cell(row=current_row+1, column=col_pos).alignment = align_center
                ws.cell(row=current_row+2, column=col_pos, value="DNI / Firma Digital / Huella").font = Font(name="Arial", size=8, italic=True)
                ws.cell(row=current_row+2, column=col_pos).alignment = align_center
    
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2, 3]:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 18)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# --- INTERFAZ COMPLETA DE STREAMLIT ---

st.markdown('<div class="main-title">🗳️ Sistema de Control y Formatos Electorales</div>', unsafe_allow_html=True)
st.markdown('<div class="subtitle">Oficina General de Coordinación Electoral (GOECOR)</div>', unsafe_allow_html=True)

# --- PANEL DE NAVEGACIÓN ---
st.sidebar.header("Panel de Navegación")

tipo_linea = st.sidebar.selectbox(
    "Seleccione la Línea de Formato:",
    options=["FM - JEL", "FM - CAE"]
)

# Renderizado condicional
if tipo_linea == "FM - JEL":
    formato_seleccionado = st.sidebar.selectbox(
        "Seleccione el Formato Oficial:",
        options=list(FORMATOS_JEL.keys()),
        format_func=lambda x: f"{FORMATOS_JEL[x]['codigo']}"
    )
    info_f = FORMATOS_JEL[formato_seleccionado]
    nombre_archivo_fisico = NOMBRES_EXACTOS_JEL[formato_seleccionado]
else:
    formato_seleccionado = st.sidebar.selectbox(
        "Seleccione el Formato Oficial:",
        options=list(FORMATOS_CAE.keys()),
        format_func=lambda x: f"{FORMATOS_CAE[x]['codigo']}"
    )
    info_f = FORMATOS_CAE[formato_seleccionado]
    nombre_archivo_fisico = NOMBRES_EXACTOS_CAE[formato_seleccionado]


# --- 📸 DIGITALIZACIÓN DE FORMATOS Y GUARDADO VÍA WEBHOOK (GOOGLE APPS SCRIPT) ---
with st.sidebar:
    st.markdown("---")
    st.markdown("### 📸 Digitalización de Formatos")
    
    if tipo_linea == "FM - JEL":
        lista_opciones_subida = [FORMATOS_JEL[k]["codigo"] for k in FORMATOS_JEL.keys()]
    else:
        lista_opciones_subida = [FORMATOS_CAE[k]["codigo"] for k in FORMATOS_CAE.keys()]
        
    formato_a_subir = st.selectbox(
        "¿Qué formato vas a subir o fotografiar?",
        options=lista_opciones_subida
    )
    
    nombre_persona = st.text_input(
        "✍️ Nombre de la Persona / Responsable:",
        placeholder="Ej. Juan Perez",
        help="El nombre que coloques aquí se usará directamente para guardar los archivos.",
        key=f"nombre_ready_{st.session_state.widget_counter}"
    )
    
    if st.session_state.mensaje_exito:
        st.success(st.session_state.mensaje_exito)
        st.session_state.mensaje_exito = None

    uploader_dinamico_key = f"uploader_ready_{st.session_state.widget_counter}"
    camera_dinamica_key = f"camera_ready_{st.session_state.widget_counter}"

    archivos_subidos = st.file_uploader(
        "SUBIR FOTO DE TUS FM AQUÍ 📤", 
        type=['png', 'jpg', 'jpeg'], 
        accept_multiple_files=True,
        key=uploader_dinamico_key
    )
    
    activar_camara = st.checkbox("📷 Activar cámara para tomar foto")
    foto_camara = None
    
    if activar_camara:
        foto_camara = st.camera_input("Captura tu formato aquí", key=camera_dinamica_key)
    
    # --- ☁️ LÓGICA DE ALMACENAMIENTO VÍA GOOGLE APPS SCRIPT ---
    if archivos_subidos or foto_camara:
        
        # URL EXACTA DEL SCRIPT DE GOOGLE PROPORCIONADA POR EL USUARIO
        WEB_APP_URL = "https://script.google.com/macros/s/AKfycbweOZOmH5MPe_iTZbUjKoF7oDQPr8t-15Pcf07w7GtNuUnVPnlGe1AML5_1ii91EwRfEA/exec"
        
        codigo_limpio = formato_a_subir.replace('/', '-')
        nombre_carpeta = f"FOTOS DE {codigo_limpio}"
        
        if nombre_persona.strip():
            nombre_limpio_archivo = "".join(c for c in nombre_persona if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
        else:
            nombre_limpio_archivo = "ANONIMO"
        
        timestamp_actual = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        contador_guardados = 0
        
        with st.spinner("Subiendo a Google Drive de forma segura..."):
            try:
                # Subir múltiples fotos de archivo
                if archivos_subidos:
                    for i, archivo_img in enumerate(archivos_subidos):
                        nombre_archivo_final = f"{codigo_limpio}_{nombre_limpio_archivo}_UPLOAD_{timestamp_actual}_{i}.png"
                        
                        # Convertir a Base64 para enviar por internet
                        foto_bytes = archivo_img.getvalue()
                        b64_img = base64.b64encode(foto_bytes).decode('utf-8')
                        
                        payload = {
                            "folderName": nombre_carpeta,
                            "filename": nombre_archivo_final,
                            "mimetype": "image/png",
                            "base64": b64_img
                        }
                        
                        response = requests.post(WEB_APP_URL, json=payload)
                        if response.status_code == 200 and response.json().get("status") == "success":
                            contador_guardados += 1
                        else:
                            st.error(f"Error al subir: {response.text}")
                        
                # Subir foto capturada con cámara web
                if foto_camara:
                    nombre_archivo_final = f"{codigo_limpio}_{nombre_limpio_archivo}_CAM_{timestamp_actual}.png"
                    foto_bytes = foto_camara.getvalue()
                    b64_img = base64.b64encode(foto_bytes).decode('utf-8')
                    
                    payload = {
                        "folderName": nombre_carpeta,
                        "filename": nombre_archivo_final,
                        "mimetype": "image/png",
                        "base64": b64_img
                    }
                    
                    response = requests.post(WEB_APP_URL, json=payload)
                    if response.status_code == 200 and response.json().get("status") == "success":
                        contador_guardados += 1
                    else:
                        st.error(f"Error al subir: {response.text}")
                    
                if contador_guardados > 0:
                    st.session_state.mensaje_exito = f"💥 ¡SUBIDO EXITOSO!\n\nSe han guardado **{contador_guardados}** archivo(s) de **`{nombre_limpio_archivo}`** directamente en tu ☁️ **Google Drive**."
                    st.session_state.widget_counter += 1  
                    st.rerun()  

            except Exception as e:
                st.error(f"❌ Fallo en la conexión con Apps Script: {e}")

# --- RENDERIZADO PRINCIPAL ---
col_izq, col_der = st.columns([2, 3])

with col_izq:
    st.markdown("### 📋 Especificaciones Técnicas")
    st.markdown(f"""
    <div class="card" style="color: #000000;">
        <h4><b style="color: #1F4E79;">Código Oficial:</b> {info_f['codigo']}</h4>
        <p><b>Denominación:</b> {info_f['nombre']}</p>
        <p><b>Propósito General:</b> {info_f['proposito']}</p>
    </div>
    """, unsafe_allow_html=True)
    
    st.markdown("### ✍️ Protocolo de Firmas")
    st.markdown(f'<div class="signature-box">🔒 {info_f["firmas"]}</div>', unsafe_allow_html=True)
    
    st.markdown("### 📥 Descarga de Plantilla Oficial")
    st.write(f"De acuerdo al flujo del manual, se sugerir el uso de formato **{info_f['tipo_descarga']}**.")
    
    # Identificar carpeta y archivo físico
    carpeta_actual = os.path.dirname(os.path.abspath(__file__))
    ruta_archivo = os.path.join(carpeta_actual, nombre_archivo_fisico)
    
    if os.path.exists(ruta_archivo):
        with open(ruta_archivo, "rb") as file:
            st.download_button(
                label=f"💾 Descargar {info_f['codigo']} (Archivo Original)",
                data=file,
                file_name=nombre_archivo_fisico,
                mime="application/octet-stream",
                key=f"dl_real_{tipo_linea}_{formato_seleccionado}"
            )
        st.success(f"✅ Encontrado: {nombre_archivo_fisico}")
    else:
        st.warning(f"⚠️ No se encontró en la carpeta el archivo: {nombre_archivo_fisico}")
        st.info("💡 Ejecutando generador dinámico de respaldo en tiempo real.")
        
        excel_data = crear_excel_formato(formato_seleccionado, info_f)
        st.download_button(
            label=f"⚙️ Generar y Descargar {info_f['codigo']} (.xlsx)",
            data=excel_data,
            file_name=f"{info_f['codigo'].replace('/', '_')}_Generado.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"dl_gen_{tipo_linea}_{formato_seleccionado}"
        )

with col_der:
    st.markdown(f"### 🔍 Vista Previa Estructurada de Datos")
    st.write("Estructura de la matriz de datos fiel al documento original:")
    
    df_preview = pd.DataFrame(info_f["datos_ejemplo"], columns=info_f["columnas"])
    st.dataframe(df_preview, width='stretch', hide_index=True)
    
    st.markdown("#### 📊 Métricas Operativas Asociadas al Formato")
    m1, m2 = st.columns(2)
    
    if tipo_linea == "FM - JEL":
        if formato_seleccionado in ["FM02", "FM11"]:
            m1.metric("Momento de Aplicación", "Pre y Post Jornada")
            m2.metric("Nivel de Criticidad", "Alto (Infraestructura)")
        elif formato_seleccionado in ["FM03", "FM05"]:
            m1.metric("Momento de Aplicación", "Mañana (Instalación)")
            m2.metric("Nivel de Criticidad", "Crítico (SIDE)")
        elif formato_seleccionado in ["FM07", "FM12", "FM13", "FM16"]:
            m1.metric("Momento de Aplicación", "Tarde / Cierre")
            m2.metric("Nivel de Criticidad", "Máximo (Actas)")
        else:
            m1.metric("Momento de Aplicación", "Durante el Día D")
            m2.metric("Nivel de Criticidad", "Medio")
    else:
        m1.metric("Momento de Aplicación", "Etapa de Capacitación")
        if formato_seleccionado in ["FM01", "FM03", "FM06"]:
            m2.metric("Nivel de Criticidad", "Alto (Miembros de Mesa)")
        elif formato_seleccionado == "FM04":
            m2.metric("Nivel de Criticidad", "Medio (Personeros)")
        elif formato_seleccionado == "FM05":
            m2.metric("Nivel de Criticidad", "Alto (Fuerzas del Orden)")
        else:
            m2.metric("Nivel de Criticidad", "Medio (Ciudadanos)")

st.markdown("---")
st.markdown("### 📑 Resumen General del Catálogo GOECOR")
with st.expander("Ver lista resumida de responsabilidades de firmas"):
    tabla_resumen = []
    for k, v in FORMATOS_JEL.items():
        tabla_resumen.append({"Línea": "FM - JEL", "Código": v["codigo"], "Nombre": v["nombre"], "Quién Firma": v["firmas"]})
    for k, v in FORMATOS_CAE.items():
        tabla_resumen.append({"Línea": "FM - CAE", "Código": v["codigo"], "Nombre": v["nombre"], "Quién Firma": v["firmas"]})
    st.table(pd.DataFrame(tabla_resumen))
